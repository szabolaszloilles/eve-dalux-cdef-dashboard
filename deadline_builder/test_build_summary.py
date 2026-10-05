"""Regression coverage for source selection, visibility and header layouts."""

from io import BytesIO
import unittest
from datetime import date

from openpyxl import Workbook

from build_summary import choose_records, read_source


class HiddenRowTests(unittest.TestCase):
    HEADERS = ['No.', 'Subject', 'Deadline', 'Responsible', 'Status', 'Work Package']

    def read_workbook(self, sheets):
        """Save real worksheet visibility settings before invoking the reader."""
        workbook = Workbook()
        workbook.remove(workbook.active)
        for name, rows, hidden_rows in sheets:
            sheet = workbook.create_sheet(name)
            sheet.append(self.HEADERS)
            for row in rows:
                sheet.append(row)
            for row_number in hidden_rows:
                sheet.row_dimensions[row_number].hidden = True
        with BytesIO() as source:
            workbook.save(source)
            workbook.close()
            source.seek(0)
            return read_source(source, 2026, 'tens')

    def test_data_only_cdef_does_not_enter_contractor_summary(self):
        entries, _ = self.read_workbook([
            ('Data', [
                ['CDEF1366', 'Improper steel cutting.', date(2026, 7, 31), 'Export owner', 'Ongoing', 'CREC MEP'],
            ], []),
            ('CREC MEP', [['CDEF1277', 'Listed CDEF', date(2026, 9, 25)]], []),
        ])

        selected = choose_records(entries, 'contractor')

        self.assertEqual([entry['id'] for entry in selected], ['CDEF1277'])
        self.assertEqual(selected[0]['company'], 'CREC MEP')
        self.assertEqual(selected[0]['date'], date(2026, 9, 25))

    def test_data_can_enrich_existing_cdef_but_not_replace_its_deadline(self):
        entries, _ = self.read_workbook([
            ('Data', [
                ['CDEF1', 'Export subject', date(2026, 7, 31), 'Export owner', 'Ongoing', 'CREC MEP'],
                ['CDEF2', 'Export-only deadline', date(2026, 8, 1)],
            ], []),
            ('CREC MEP', [
                ['CDEF1', 'Listed subject', date(2026, 9, 30)],
                ['CDEF2', 'No contractor deadline', None],
            ], []),
        ])

        selected = choose_records(entries, 'contractor')

        self.assertEqual(len(selected), 1)
        self.assertEqual(selected[0]['id'], 'CDEF1')
        self.assertEqual(selected[0]['date'], date(2026, 9, 30))
        self.assertEqual(selected[0]['display_subject'], 'Listed subject')
        self.assertEqual(selected[0]['display_owner'], 'Export owner')
        self.assertEqual(selected[0]['display_status'], 'Ongoing')

    def test_data_only_workbook_yields_no_contractor_records(self):
        entries, _ = self.read_workbook([
            ('Data', [['CDEF1366', 'Export only', date(2026, 7, 31)]], []),
        ])

        self.assertEqual(choose_records(entries, 'contractor'), [])
        self.assertEqual(len(choose_records(entries, 'all')), 1)

    def test_hidden_earlier_deadline_and_status_cannot_affect_visible_record(self):
        entries, _ = self.read_workbook([
            ('Contractor', [
                ['CDEF1', 'Hidden subject', date(2026, 8, 1), 'Hidden owner', 'Overdue'],
                ['CDEF1', 'Visible subject', date(2026, 9, 30)],
            ], [2]),
        ])

        self.assertEqual([entry['hidden'] for entry in entries], [True, False])
        selected = choose_records(entries, 'contractor')

        self.assertEqual(len(selected), 1)
        record = selected[0]
        self.assertEqual(record['row'], 3)
        self.assertEqual(record['date'], date(2026, 9, 30))
        self.assertEqual(record['display_subject'], 'Visible subject')
        self.assertEqual(record['display_owner'], 'Not supplied')
        self.assertEqual(record['display_status'], 'Not supplied')
        self.assertEqual(record['other_deadlines'], '')
        self.assertEqual(record['selection_note'], '')

    def test_hidden_contractor_prevents_data_fallback_even_without_deadline(self):
        for contractor_deadline in (date(2026, 8, 1), None):
            with self.subTest(contractor_deadline=contractor_deadline):
                entries, _ = self.read_workbook([
                    ('Data', [['CDEF1', 'Export subject', date(2026, 9, 30)]], []),
                    ('Contractor', [['CDEF1', 'Hidden subject', contractor_deadline]], [2]),
                ])

                self.assertEqual(choose_records(entries, 'contractor'), [])

    def test_visible_copy_on_another_contractor_sheet_remains_eligible(self):
        entries, _ = self.read_workbook([
            ('Data', [['CDEF1', 'Export subject', date(2026, 8, 1)]], []),
            ('First contractor', [['CDEF1', 'Hidden subject', date(2026, 8, 2)]], [2]),
            ('Second contractor', [['CDEF1', 'Visible subject', date(2026, 9, 30)]], []),
        ])

        selected = choose_records(entries, 'contractor')

        self.assertEqual(len(selected), 1)
        self.assertEqual(selected[0]['sheet'], 'Second contractor')
        self.assertEqual(selected[0]['date'], date(2026, 9, 30))
        self.assertNotIn('First contractor', selected[0]['other_deadlines'])

    def test_different_contractor_duplicates_are_flagged(self):
        entries, _ = self.read_workbook([
            ('CSCEC5 MEP', [['CDEF872', 'Damaged panel', date(2026, 11, 10)]], []),
            ('Huake', [['CDEF872', 'Damaged panel', date(2026, 12, 31)]], []),
        ])
        selected = choose_records(entries, 'contractor')
        self.assertEqual(len(selected), 1)
        self.assertEqual(selected[0]['sheet'], 'CSCEC5 MEP')
        self.assertTrue(selected[0]['assignment_review'])
        self.assertIn('Also listed in Huake', selected[0]['selection_note'])
        self.assertIn('Huake!', selected[0]['other_deadlines'])

    def test_hidden_data_row_cannot_enrich_visible_contractor_metadata(self):
        entries, _ = self.read_workbook([
            ('Data', [
                ['CDEF1', 'Hidden subject', date(2026, 8, 1), 'Hidden owner', 'Hidden status'],
            ], [2]),
            ('Contractor', [['CDEF1', None, date(2026, 9, 30)]], []),
        ])

        selected = choose_records(entries, 'contractor')

        self.assertEqual(len(selected), 1)
        for field in ('subject', 'owner', 'status'):
            self.assertEqual(selected[0][f'display_{field}'], 'Not supplied')
            self.assertEqual(selected[0][f'{field}_source'], '')
        self.assertEqual(selected[0]['other_deadlines'], '')
        self.assertEqual(selected[0]['selection_note'], '')

    def test_all_policy_includes_visible_data_but_excludes_hidden_source_rows(self):
        entries, _ = self.read_workbook([
            ('Data', [
                ['CDEF1', 'Visible export', date(2026, 9, 30)],
                ['CDEF2', 'Hidden export', date(2026, 8, 1)],
            ], [3]),
            ('Contractor', [
                ['CDEF1', 'Hidden contractor', date(2026, 8, 2)],
                ['CDEF2', 'Visible contractor', date(2026, 9, 20)],
            ], [2]),
        ])

        selected = choose_records(entries, 'all')

        self.assertEqual(
            {(entry['id'], entry['sheet']) for entry in selected},
            {('CDEF1', 'Data'), ('CDEF2', 'Contractor')},
        )
        self.assertTrue(all(not entry['hidden'] for entry in selected))
        self.assertTrue(all(entry['other_deadlines'] == '' for entry in selected))

    def test_all_hidden_source_yields_no_selected_records(self):
        entries, _ = self.read_workbook([
            ('Data', [['CDEF1', 'Hidden export', date(2026, 8, 1)]], [2]),
            ('Contractor', [['CDEF1', 'Hidden contractor', date(2026, 9, 30)]], [2]),
        ])

        for policy in ('contractor', 'all'):
            with self.subTest(policy=policy):
                self.assertEqual(choose_records(entries, policy), [])


class HeaderLayoutTests(unittest.TestCase):
    def read_rows(self, rows, hidden_rows=()):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Huake'
        for row in rows:
            sheet.append(row)
        for row_number in hidden_rows:
            sheet.row_dimensions[row_number].hidden = True
        with BytesIO() as source:
            workbook.save(source)
            workbook.close()
            source.seek(0)
            return read_source(source, 2026, 'tens')

    def test_bilingual_correction_deadline_at_row_five(self):
        entries, inventory = self.read_rows([
            ['Project'], ['Type'], ['Prepare date'], [],
            ['No. / 编号', 'Discipline', 'Type', 'Date modified', 'Date created',
             'Building', 'Subject / 主题', 'Message', ' Correction Deadline / 纠正截止日期 ',
             'Deadline2', 'Construction Deadline'],
            ['CDEF1000', None, None, None, None, None, 'Fire barrier closure', None,
             date(2026, 10, 20), date(2026, 7, 1), date(2026, 8, 1)],
            ['CDEF1014', None, None, None, None, None, 'Hidden defect', None,
             date(2026, 9, 10)],
            ['CDEF1018', None, None, None, None, None, 'No correction deadline', None,
             None, date(2026, 8, 2)],
        ], hidden_rows=[7])

        selected = choose_records(entries, 'contractor')
        self.assertEqual(len(selected), 1)
        self.assertEqual(selected[0]['id'], 'CDEF1000')
        self.assertEqual(selected[0]['cell'], 'I6')
        self.assertEqual(selected[0]['date'], date(2026, 10, 20))
        self.assertEqual(selected[0]['display_subject'], 'Fire barrier closure')
        self.assertEqual(inventory[0][1], 'I5')

    def test_bilingual_plain_deadline_and_metadata(self):
        entries, _ = self.read_rows([
            ['No. / 编号', 'Subject / 主题', 'Deadline / 截止日期',
             'Responsible / 负责人', 'Status / 状态'],
            ['CDEF1', 'Visible subject', '2026.10.28', 'Owner', 'Ongoing'],
        ])
        selected = choose_records(entries, 'contractor')
        self.assertEqual(len(selected), 1)
        self.assertEqual(selected[0]['date'], date(2026, 10, 28))
        self.assertEqual(selected[0]['display_owner'], 'Owner')
        self.assertEqual(selected[0]['display_status'], 'Ongoing')

    def test_other_deadline_headers_are_not_substituted(self):
        entries, inventory = self.read_rows([
            ['No. / 编号', 'Deadline2 / 截止日期', 'Construction Deadline / 施工截止日期'],
            ['CDEF1', date(2026, 7, 1), date(2026, 8, 1)],
        ])
        self.assertEqual(entries, [])
        self.assertEqual(inventory[0][5], 'Excluded: no Deadline header')

    def test_two_supported_deadline_columns_require_review(self):
        with self.assertRaisesRegex(ValueError, 'Multiple Deadline headers'):
            self.read_rows([
                ['No. / 编号', 'Deadline', 'Correction Deadline / 纠正截止日期'],
                ['CDEF1', date(2026, 7, 1), date(2026, 10, 20)],
            ])


if __name__ == '__main__':
    unittest.main()
