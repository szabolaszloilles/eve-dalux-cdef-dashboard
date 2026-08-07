"""Generate the CDEF Excel report. Usable as a CLI or imported by the app."""
import sys
import io
import datetime as dt
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
import cdef_data as cd

NAVY = "1F3864"
BLUE = "2E5496"
LIGHT = "D9E1F2"
GREY = "F2F2F2"
GREEN = "C6EFCE"
GREEN_T = "006100"
RED = "FFC7CE"
RED_T = "9C0006"
WHITE = "FFFFFF"
FONT = "Calibri"

thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def hdr(cell, text, bg=NAVY, fg=WHITE, size=11, align="center"):
    cell.value = text
    cell.font = Font(name=FONT, bold=True, color=fg, size=size)
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    cell.border = border


def cellfmt(cell, value, numfmt=None, bold=False, bg=None, align="center"):
    cell.value = value
    cell.font = Font(name=FONT, bold=bold)
    cell.alignment = Alignment(horizontal=align, vertical="center")
    cell.border = border
    if numfmt:
        cell.number_format = numfmt
    if bg:
        cell.fill = PatternFill("solid", fgColor=bg)


def write_table(ws, df_, start_row=1, pct_cols=(), delta_cols=(), int_cols=(),
                freeze=False):
    ws.sheet_view.showGridLines = False
    cols = list(df_.columns)
    for j, name in enumerate(cols, start=1):
        hdr(ws.cell(row=start_row, column=j), str(name))
    for i, (_, row) in enumerate(df_.iterrows(), start=start_row + 1):
        for j, name in enumerate(cols, start=1):
            val = row[name]
            numfmt = None
            if name in pct_cols:
                numfmt = '0.0"%";(0.0)"%";"–"'
            elif name in delta_cols:
                numfmt = '+#,##0;-#,##0;"–"'
            elif name in int_cols:
                numfmt = '#,##0'
            try:
                if val is pd.NA or pd.isna(val):
                    val = None
            except (TypeError, ValueError):
                pass
            if isinstance(val, dt.date):
                val = val.strftime("%Y-%m-%d")
            align = "left" if j == 1 and not isinstance(val, (int, float)) else "center"
            cellfmt(ws.cell(row=i, column=j), val, numfmt=numfmt, align=align)
        if (i - start_row) % 2 == 0:
            for j in range(1, len(cols) + 1):
                if ws.cell(row=i, column=j).fill.fgColor.rgb in (None, "00000000"):
                    ws.cell(row=i, column=j).fill = PatternFill("solid", fgColor=GREY)
    for j, name in enumerate(cols, start=1):
        width = max(12, min(40, max([len(str(name))] + [len(str(v)) for v in df_[name].astype(str)]) + 2))
        ws.column_dimensions[get_column_letter(j)].width = width
    if freeze:
        ws.freeze_panes = ws.cell(row=start_row + 1, column=2)
    return start_row + len(df_)


def color_deltas(ws, df_, start_row, delta_cols):
    cols = list(df_.columns)
    n = len(df_)
    for name in delta_cols:
        j = cols.index(name) + 1
        L = get_column_letter(j)
        rng = f"{L}{start_row+1}:{L}{start_row+n}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], font=Font(color=GREEN_T)))
        ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], font=Font(color=RED_T)))


def _add_sheets(wb, df, top10=None, prefix="", title="CDEF — Construction Defects Dashboard"):
    """Add a full set of report sheets for `df` to an existing workbook.

    `prefix` is prepended to every sheet name so a second dataset (e.g. CDEF-SG)
    can be appended to the same workbook without clashing.
    """
    weekly = cd.rolling_metrics(df)
    matrix = cd.contractor_status_matrix(df)
    csum = cd.contractor_summary(df)

    # ---------------- Summary ----------------
    if prefix:
        ws = wb.create_sheet(f"{prefix}Summary")
    else:
        ws = wb.active
        ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:F1")
    hdr(ws["A1"], title, bg=NAVY, size=16, align="left")
    ws.row_dimensions[1].height = 28
    ws["A2"] = f"EVE Factory Project Debrecen   ·   Generated {dt.date.today():%d %b %Y}"
    ws["A2"].font = Font(name=FONT, italic=True, color="595959")
    ws.merge_cells("A2:F2")

    last = weekly.iloc[-1]
    rates = cd.week_resolution_rates(df)
    kpis = [
        ("Total defects", int(last["Total defects (cumulative)"]), int(last["Total Δ vs prev"])),
        ("New (latest week)", int(last["New defects"]), int(last["New Δ vs prev"])),
        ("Accepted (latest week)", int(last["Accepted defects"]), int(last["Accepted Δ vs prev"])),
        ("Open defects", int((~df["is_accepted"]).sum() - df[df['status'].isin(['Discontinued', 'Rejected'])].shape[0]), None),
        ("Reported ready · CÉH", cd.reported_ready_ceh(df), None),
        ("Resolved of week intake", f"{rates['pct_todate']:.0f}%", None),
    ]
    r = 4
    ws[f"A{r}"] = f"Latest reporting week (Fri–Thu): {last['ISO week']}  (ending {last['Week starting'] + dt.timedelta(days=6):%d %b %Y})"
    ws[f"A{r}"].font = Font(name=FONT, bold=True, size=12, color=NAVY)
    ws.merge_cells(f"A{r}:F{r}")
    r = 6
    col = 1
    for label, val, delta in kpis:
        c = ws.cell(row=r, column=col)
        hdr(c, label, bg=BLUE)
        v = ws.cell(row=r + 1, column=col)
        cellfmt(v, val, bold=True)
        v.font = Font(name=FONT, bold=True, size=18, color=NAVY)
        d = ws.cell(row=r + 2, column=col)
        if label.startswith("Resolved"):
            cellfmt(d, f"{rates['pct_inweek']:.0f}% within the week")
            d.font = Font(name=FONT, bold=True, color="595959")
        elif delta is None:
            cellfmt(d, "")
        else:
            arrow = "▲" if delta > 0 else ("▼" if delta < 0 else "▬")
            cellfmt(d, f"{arrow} {delta:+d} vs prev week")
            d.font = Font(name=FONT, bold=True, color=(GREEN_T if delta > 0 else RED_T) if delta else "595959")
        col += 1
    for cc in range(1, 7):
        ws.column_dimensions[get_column_letter(cc)].width = 20
    ws.row_dimensions[7].height = 30

    note_r = 11
    notes = [
        "Definitions:",
        "• Contractor = Work Package (Dalux column I).",
        "• Reporting weeks run Friday 00:00 – Thursday 24:00. Figures show the last complete Fri–Thu week.",
        "• New defects = count by date created, within each Fri–Thu week.",
        "• Accepted defects = status 'Approved' / 'Approved, follow-up', dated by last-modified date, within each Fri–Thu week.",
        "• Total defects = cumulative count of all defects raised up to and including the end of that week.",
        "• Reported ready · CÉH = defects with status 'Reported ready' (a.k.a. 'Awaiting approval') whose Role starts with 'CÉH' (awaiting CÉH approval).",
        "• Resolved of week intake = of the defects RAISED in the latest week, the share now Approved (or Approved, follow-up).",
        "   The sub-line shows the share approved within that same Fri–Thu week.",
        "• Δ vs prev = change against the previous Fri–Thu week.",
    ]
    for i, t in enumerate(notes):
        c = ws.cell(row=note_r + i, column=1)
        c.value = t
        c.font = Font(name=FONT, italic=(i > 0), bold=(i == 0), color="595959")
        ws.merge_cells(start_row=note_r + i, start_column=1, end_row=note_r + i, end_column=6)

    # ---------------- By Contractor ----------------
    ws = wb.create_sheet(f"{prefix}By Contractor")
    ws.merge_cells("A1:F1")
    hdr(ws["A1"], "Defects by Contractor (Work Package) × Status", bg=NAVY, size=13, align="left")
    ws.row_dimensions[1].height = 24
    mat = matrix.reset_index()
    end = write_table(ws, mat, start_row=3, int_cols=[c for c in mat.columns if c != "contractor"])
    for j in range(1, mat.shape[1] + 1):
        ws.cell(row=end, column=j).font = Font(name=FONT, bold=True)
        ws.cell(row=end, column=j).fill = PatternFill("solid", fgColor=LIGHT)

    ws.cell(row=end + 2, column=1, value="Contractor summary").font = Font(name=FONT, bold=True, size=12, color=NAVY)
    end2 = write_table(ws, csum, start_row=end + 3, pct_cols=["Acceptance %"], int_cols=["Total", "Accepted", "Open"])

    ctype = cd.contractor_type_summary(df)
    if not ctype.empty:
        ws.cell(row=end2 + 2, column=1,
                value="Contractor × defect type (all defects)").font = Font(
                    name=FONT, bold=True, size=12, color=NAVY)
        write_table(ws, ctype, start_row=end2 + 3,
                    pct_cols=["Resolved %"],
                    int_cols=["Total", "Resolved", "Open"])

    # ---------------- 7-day rolling sheets ----------------
    specs = [
        ("7-Day New", ["Week starting", "ISO week", "New defects", "New Δ vs prev", "New Δ %"], "New defects", "New Δ vs prev", "New Δ %"),
        ("7-Day Accepted", ["Week starting", "ISO week", "Accepted defects", "Accepted Δ vs prev", "Accepted Δ %"], "Accepted defects", "Accepted Δ vs prev", "Accepted Δ %"),
        ("7-Day Total", ["Week starting", "ISO week", "Total defects (cumulative)", "Total Δ vs prev", "Total Δ %"], "Total defects (cumulative)", "Total Δ vs prev", "Total Δ %"),
    ]
    for sheet_title, cols, valcol, dcol, pcol in specs:
        ws = wb.create_sheet(f"{prefix}{sheet_title}")
        ws.merge_cells("A1:E1")
        hdr(ws["A1"], sheet_title.replace("7-Day", "7-Day rolling —"), bg=NAVY, size=13, align="left")
        ws.row_dimensions[1].height = 24
        sub = weekly[cols].copy()
        sub = sub.rename(columns={"Week starting": "Period start", "ISO week": "7-day period"})
        start = 3
        end = write_table(ws, sub, start_row=start, pct_cols=[pcol], delta_cols=[dcol], int_cols=[valcol], freeze=True)
        color_deltas(ws, sub, start, [dcol])

    # ---------------- Raw CDEF ----------------
    # ---------------- Resolution by defect type ----------------
    if rates["by_type"]:
        ws = wb.create_sheet(f"{prefix}By Defect Type")
        ws.merge_cells("A1:F1")
        hdr(ws["A1"], f"Resolution by defect type — {rates['label']} intake",
            bg=NAVY, size=13, align="left")
        ws.row_dimensions[1].height = 24
        ws["A2"] = (f"Of the {rates['new']} defects raised in the reporting week, "
                    f"how many are now resolved (Approved / Approved, follow-up).")
        ws["A2"].font = Font(name=FONT, italic=True, color="595959")
        ws.merge_cells("A2:F2")
        bt = pd.DataFrame(rates["by_type"])[
            ["type", "new", "resolved_todate", "pct_todate", "resolved_inweek", "pct_inweek"]]
        bt.columns = ["Defect type", "Raised", "Resolved (to date)", "Resolved %",
                      "Resolved in week", "In-week %"]
        write_table(ws, bt, start_row=4,
                    pct_cols=["Resolved %", "In-week %"],
                    int_cols=["Raised", "Resolved (to date)", "Resolved in week"])

    # ---------------- Resolution by contractor (x defect type) ----------------
    if rates.get("by_contractor"):
        ws = wb.create_sheet(f"{prefix}Weekly by Contractor")
        ws.merge_cells("A1:F1")
        hdr(ws["A1"], f"Resolution by contractor — {rates['label']} intake",
            bg=NAVY, size=13, align="left")
        ws.row_dimensions[1].height = 24
        ws["A2"] = (f"Of the defects raised against each contractor in the reporting week, "
                    f"how many are now resolved (Approved / Approved, follow-up).")
        ws["A2"].font = Font(name=FONT, italic=True, color="595959")
        ws.merge_cells("A2:F2")
        bc = pd.DataFrame(rates["by_contractor"])[
            ["contractor", "new", "resolved_todate", "pct_todate",
             "resolved_inweek", "pct_inweek"]]
        bc.columns = ["Contractor", "Raised", "Resolved (to date)", "Resolved %",
                      "Resolved in week", "In-week %"]
        _end = write_table(ws, bc, start_row=4,
                           pct_cols=["Resolved %", "In-week %"],
                           int_cols=["Raised", "Resolved (to date)", "Resolved in week"])

        if rates.get("by_contractor_type"):
            ws.cell(row=_end + 2, column=1,
                    value="Split by defect type").font = Font(
                        name=FONT, bold=True, size=12, color=NAVY)
            bct = pd.DataFrame(rates["by_contractor_type"])[
                ["contractor", "type", "new", "resolved_todate", "pct_todate",
                 "resolved_inweek", "pct_inweek"]]
            bct.columns = ["Contractor", "Defect type", "Raised", "Resolved (to date)",
                           "Resolved %", "Resolved in week", "In-week %"]
            bct = bct.sort_values(["Contractor", "Raised"], ascending=[True, False])
            write_table(ws, bct, start_row=_end + 3,
                        pct_cols=["Resolved %", "In-week %"],
                        int_cols=["Raised", "Resolved (to date)", "Resolved in week"])

    # ---------------- Contractor performance overview ----------------
    perf = []
    wk_map = {c["contractor"]: c for c in rates.get("by_contractor", [])}
    for _, row in csum.iterrows():
        c = row["contractor"]
        t10row = None
        if top10 is not None and not top10.empty:
            m = top10[top10["contractor"] == c]
            if not m.empty:
                t10row = m.iloc[0]
        if t10row is None and c not in wk_map and row["Total"] < 5:
            continue
        perf.append({
            "Contractor": c,
            "Top 10 reply %": (round(float(t10row["reply_pct"]), 1)
                               if t10row is not None else None),
            "Week closing %": (round(wk_map[c]["pct_todate"], 1)
                               if c in wk_map else None),
            "Overall closing %": float(row["Acceptance %"]),
        })
    if perf:
        ws = wb.create_sheet(f"{prefix}Contractor Performance")
        ws.merge_cells("A1:E1")
        hdr(ws["A1"], f"Contractor performance — {rates['label']}", bg=NAVY,
            size=13, align="left")
        ws.row_dimensions[1].height = 24
        ws["A2"] = ("Top 10 reply % = tracker issues marked approved/resolved. "
                    "Week closing % = defects raised in the reporting week now resolved. "
                    "Overall closing % = all defects resolved to date.")
        ws["A2"].font = Font(name=FONT, italic=True, color="595959")
        ws.merge_cells("A2:E2")
        write_table(ws, pd.DataFrame(perf), start_row=4,
                    pct_cols=["Top 10 reply %", "Week closing %", "Overall closing %"])

    ws = wb.create_sheet(f"{prefix}Raw CDEF")
    raw = df[["id", "subject", "contractor", "defectType", "status", "statusDetailed",
              "discipline", "role", "responsibleCompany", "created", "modified",
              "is_accepted"]].copy()
    raw["created"] = raw["created"].apply(lambda d: d.strftime("%Y-%m-%d %H:%M") if pd.notna(d) else "")
    raw["modified"] = raw["modified"].apply(lambda d: d.strftime("%Y-%m-%d %H:%M") if pd.notna(d) else "")
    raw.columns = ["ID", "Subject", "Contractor", "Defect type", "Status", "Status detail",
                   "Discipline", "Role", "Responsible company", "Created", "Modified",
                   "Accepted"]
    write_table(ws, raw, start_row=1, freeze=True)
    ws.auto_filter.ref = f"A1:L{len(raw)+1}"
    ws.column_dimensions["B"].width = 45

    return wb


def build_workbook(df, top10=None, df_sg=None):
    """Build the report workbook.

    `df`     - regular CDEF defects (may be None if only SG data is supplied)
    `df_sg`  - optional CDEF-SG defects; their sheets are appended to the same
               workbook with an 'SG ' prefix.
    """
    wb = Workbook()
    first_done = False
    if df is not None and not df.empty:
        _add_sheets(wb, df, top10=top10, prefix="",
                    title="CDEF — Construction Defects Dashboard")
        first_done = True
    if df_sg is not None and not df_sg.empty:
        if not first_done:
            # SG only: use the default sheet for its Summary
            _add_sheets(wb, df_sg, top10=top10, prefix="",
                        title="CDEF-SG — Construction Defects Dashboard")
        else:
            _add_sheets(wb, df_sg, top10=top10, prefix="SG ",
                        title="CDEF-SG — Construction Defects Dashboard")
    return wb


def build_report_bytes(df, top10=None, df_sg=None):
    """Return the report workbook as .xlsx bytes (for in-app download)."""
    wb = build_workbook(df, top10=top10, df_sg=df_sg)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


if __name__ == "__main__":
    in_path = sys.argv[1] if len(sys.argv) > 1 else "Reports.xlsx"
    out = sys.argv[2] if len(sys.argv) > 2 else "CDEF_Report.xlsx"
    df = cd.load_cdef_any(in_path)
    build_workbook(df).save(out)
    print("saved", out)
